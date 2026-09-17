import pytest

import config
from ingestion import DocumentError, chunk_document, hash_file, load_document


def test_txt_loading_and_chunking(tmp_path):
    p = tmp_path / "doc.txt"
    p.write_text("Paragraph one.\n\nParagraph two is a bit longer than one.\n\n" * 40)
    pages = load_document(str(p))
    chunks = chunk_document(pages, "doc.txt")
    assert len(chunks) > 1
    assert all(len(c["text"]) <= config.CHUNK_SIZE + 2 * config.CHUNK_OVERLAP for c in chunks)
    assert all(c["source"] == "doc.txt" for c in chunks)


def test_md_loading(tmp_path):
    p = tmp_path / "doc.md"
    p.write_text("# Title\n\nSome content.")
    pages = load_document(str(p))
    assert "Title" in pages[0]["text"]
    assert pages[0]["page"] is None


def test_docx_loading(tmp_path):
    from docx import Document

    p = tmp_path / "doc.docx"
    doc = Document()
    doc.add_paragraph("Hello from docx.")
    doc.save(str(p))

    pages = load_document(str(p))
    assert "Hello from docx" in pages[0]["text"]


def test_unsupported_extension_rejected(tmp_path):
    p = tmp_path / "doc.xyz"
    p.write_text("nope")
    with pytest.raises(DocumentError):
        load_document(str(p))


def test_oversized_file_rejected(tmp_path, monkeypatch):
    monkeypatch.setattr(config, "MAX_FILE_SIZE_MB", 0.0001)
    p = tmp_path / "big.txt"
    p.write_text("x" * 10_000)
    with pytest.raises(DocumentError):
        load_document(str(p))


def test_scanned_pdf_detected(tmp_path):
    from pypdf import PdfWriter

    p = tmp_path / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)  # no text stream at all
    with open(p, "wb") as f:
        writer.write(f)

    with pytest.raises(DocumentError):
        load_document(str(p))


def test_text_pdf_extracts_with_page_numbers(tmp_path):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("Helvetica", size=12)
    pdf.cell(0, 10, "Employees receive 20 days of PTO per year.")
    p = tmp_path / "handbook.pdf"
    pdf.output(str(p))

    pages = load_document(str(p))
    assert pages[0]["page"] == 1
    assert "PTO" in pages[0]["text"]


def test_hash_file_is_deterministic_and_content_sensitive(tmp_path):
    a = tmp_path / "a.txt"
    a.write_text("content")
    b = tmp_path / "b.txt"
    b.write_text("content")
    c = tmp_path / "c.txt"
    c.write_text("different content")

    assert hash_file(str(a)) == hash_file(str(a))
    assert hash_file(str(a)) == hash_file(str(b))  # same content, different name
    assert hash_file(str(a)) != hash_file(str(c))
